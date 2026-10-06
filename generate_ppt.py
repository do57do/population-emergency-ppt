from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor


OUT_FILE = "population_emergency_management.pptx"


def hex_to_rgb(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = hex_to_rgb(color)


def add_title(slide, title, subtitle=None, title_color="1F2D3D", subtitle_color="4A607A"):
    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.35), Inches(12), Inches(0.7))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.LEFT
    run = p.runs[0]
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = hex_to_rgb(title_color)
    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.75), Inches(1.0), Inches(11), Inches(0.4))
        s_tf = sub_box.text_frame
        p2 = s_tf.paragraphs[0]
        p2.text = subtitle
        p2.alignment = PP_ALIGN.LEFT
        run2 = p2.runs[0]
        run2.font.size = Pt(12)
        run2.font.color.rgb = hex_to_rgb(subtitle_color)


def add_bullets(slide, bullets, left=0.9, top=1.5, width=11.5, height=5.5, font_size=20, color="1F2D3D"):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 8
    tf.margin_right = 8
    tf.margin_top = 8
    tf.margin_bottom = 8
    for idx, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.bullet = True
        p.alignment = PP_ALIGN.LEFT
        run = p.runs[0]
        run.font.size = Pt(font_size)
        run.font.color.rgb = hex_to_rgb(color)
        run.font.name = "Microsoft YaHei"


def add_footer(slide, text="社会学概论 · 人口章节专题汇报"):
    footer = slide.shapes.add_textbox(Inches(0.7), Inches(6.9), Inches(12), Inches(0.3))
    tf = footer.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.RIGHT
    run = p.runs[0]
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(120, 130, 150)


def make_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, "F4F7FB")
    # background geometric blocks
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.33), Inches(0.18))
    shape.fill.solid()
    shape.fill.fore_color.rgb = hex_to_rgb("0B3D91")
    shape.line.fill.background()

    # decorative circles
    for x, y, size, color in [
        (10.5, 0.9, 3.2, "D8E7FF"),
        (11.2, 2.4, 2.5, "B7D1FF"),
        (9.0, 3.5, 1.8, "EAF1FF"),
    ]:
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(size), Inches(size))
        circle.fill.solid()
        circle.fill.fore_color.rgb = hex_to_rgb(color)
        circle.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.9), Inches(1.7), Inches(9.8), Inches(1.6))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "人口要素：社会基础与应急管理的底层逻辑"
    p.alignment = PP_ALIGN.LEFT
    run = p.runs[0]
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = hex_to_rgb("1C2B50")

    sub = slide.shapes.add_textbox(Inches(0.9), Inches(3.0), Inches(7.8), Inches(0.7))
    tf2 = sub.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = "社会学概论 · 人口章节专题汇报"
    p2.alignment = PP_ALIGN.LEFT
    run2 = p2.runs[0]
    run2.font.size = Pt(18)
    run2.font.color.rgb = hex_to_rgb("48627F")

    meta = slide.shapes.add_textbox(Inches(0.9), Inches(4.6), Inches(4.5), Inches(1.5))
    tf3 = meta.text_frame
    for i, line in enumerate(["汇报人：张三", "专业：应急管理", "时间：2026年10月"]):
        p3 = tf3.paragraphs[0] if i == 0 else tf3.add_paragraph()
        p3.text = line
        p3.alignment = PP_ALIGN.LEFT
        run3 = p3.runs[0]
        run3.font.size = Pt(18)
        run3.font.color.rgb = hex_to_rgb("1F2D3D")

    # simple population tower graphic by rectangles
    tower_x = 10.3
    tower_y = 2.4
    for h in [0.8, 1.1, 1.4, 1.8, 2.0]:
        rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(tower_x), Inches(tower_y + 2.2 - h), Inches(1.2), Inches(h))
        rect.fill.solid()
        rect.fill.fore_color.rgb = hex_to_rgb("9BC1FF")
        rect.line.fill.background()
        tower_x += 0.55

    add_footer(slide, "应急管理专业专题汇报")


def make_toc(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, "F8FAFC")
    add_title(slide, "目录", subtitle="Population × Emergency Management")
    left = 0.9
    top = 1.55
    items = [
        "1. 社会学概论「人口」章节核心理论回顾",
        "2. 应急管理专业知识体系框架",
        "3. 人口与应急管理的内在共通逻辑",
        "4. 人口各维度对应急管理五大模块的联系",
        "5. 现实案例分析",
        "6. 总结与专业思考",
    ]
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(9.5), Inches(4.5))
    tf = box.text_frame
    for idx, text in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = text
        p.level = 0
        p.bullet = True
        p.alignment = PP_ALIGN.LEFT
        run = p.runs[0]
        run.font.size = Pt(22)
        run.font.color.rgb = hex_to_rgb("1F2D3D")
    # decorative block
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10.3), Inches(1.5), Inches(2.0), Inches(4.7))
    rect.fill.solid()
    rect.fill.fore_color.rgb = hex_to_rgb("E7F0FF")
    rect.line.fill.background()
    label = slide.shapes.add_textbox(Inches(10.5), Inches(2.8), Inches(1.5), Inches(1.0))
    lab = label.text_frame
    p = lab.paragraphs[0]
    p.text = "核心问题"
    p.alignment = PP_ALIGN.CENTER
    run = p.runs[0]
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = hex_to_rgb("0B3D91")
    add_footer(slide)


def make_sociology_review(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, "F8FAFC")
    add_title(slide, "社会学概论：人口章节核心内容", subtitle="马工程《社会学概论》")
    add_bullets(
        slide,
        [
            "人口是社会存在和发展的物质基础，人口状况深刻制约社会运行。",
            "人口五大维度：规模、结构、空间分布、迁移流动、质量与健康。",
            "人口结构重点关注：年龄结构、性别结构、城乡结构、教育与职业结构。",
            "脆弱群体：老人、儿童、残障者、流动人口，需要在灾害中优先关注。",
            "人口不只是数字，而是理解社会风险与社会治理的重要切入点。",
        ],
        left=0.9,
        top=1.6,
        width=7.8,
        height=4.6,
        font_size=19,
    )
    # side panel
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.8), Inches(1.6), Inches(3.6), Inches(4.7))
    panel.fill.solid()
    panel.fill.fore_color.rgb = hex_to_rgb("EAF2FF")
    panel.line.fill.background()
    text = slide.shapes.add_textbox(Inches(9.1), Inches(2.0), Inches(3.0), Inches(3.2))
    tf = text.text_frame
    for i, line in enumerate(["1. 规模", "2. 结构", "3. 空间分布", "4. 迁移流动", "5. 质量与健康"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = PP_ALIGN.CENTER
        run = p.runs[0]
        run.font.size = Pt(18)
        run.font.color.rgb = hex_to_rgb("0B3D91")
    add_footer(slide)


def make_emergency_framework(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, "F5F7FB")
    add_title(slide, "应急管理专业知识体系框架", subtitle="国内高校应急管理本科通用五大核心模块")
    items = [
        "风险识别与风险评估：识别危险源，评估概率与损失，划分风险等级。",
        "应急预案与应急规划：编制预案、避难场所和疏散方案、防灾减灾规划。",
        "应急指挥与现场处置：现场搜救、现场管控、抢险救援与协调。",
        "应急资源管理与灾民安置：物资储备、临时安置、生活保障。",
        "灾后恢复重建与社会治理：心理干预、秩序恢复、重建和次生风险排查。",
    ]
    add_bullets(slide, items, left=0.8, top=1.6, width=12.0, height=4.5, font_size=18)
    # callout strip
    strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(5.7), Inches(12.0), Inches(0.7))
    strip.fill.solid()
    strip.fill.fore_color.rgb = hex_to_rgb("DCEBFF")
    strip.line.fill.background()
    text = slide.shapes.add_textbox(Inches(1.0), Inches(5.82), Inches(10.0), Inches(0.3))
    t = text.text_frame
    p = t.paragraphs[0]
    p.text = "补充：应急管理是交叉学科，融合公共管理、法学、社会学、安全工程等多学科知识。"
    p.alignment = PP_ALIGN.LEFT
    run = p.runs[0]
    run.font.size = Pt(12)
    run.font.color.rgb = hex_to_rgb("1F2D3D")
    add_footer(slide)


def make_common_logic(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, "F8FAFC")
    add_title(slide, "人口与应急管理的六大共通点", subtitle="核心重点")
    x_positions = [0.7, 4.6, 8.5]
    y_positions = [1.7, 3.8]
    titles = ["价值目标共通", "研究对象共通", "风险逻辑共通", "治理思维共通", "实践路径共通", "最终目标共通"]
    descs = [
        "都以人的生存发展为根本。",
        "都关注规模、结构、分布、迁移和脆弱群体。",
        "人口失衡会催生社会风险，放大灾害损失。",
        "重视全过程管理，从预判到恢复。",
        "依靠人口数据和多方协同治理。",
        "都致力于维护社会稳定与秩序。",
    ]
    for idx, (title, desc) in enumerate(zip(titles, descs)):
        row = idx // 3
        col = idx % 3
        x = x_positions[col]
        y = y_positions[row]
        rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(3.1), Inches(1.5))
        rect.fill.solid()
        rect.fill.fore_color.rgb = hex_to_rgb("EAF1FF")
        rect.line.color.rgb = hex_to_rgb("B9CBE8")
        text = slide.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.15), Inches(2.7), Inches(1.2))
        tf = text.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.alignment = PP_ALIGN.LEFT
        run = p.runs[0]
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = hex_to_rgb("0B3D91")
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.alignment = PP_ALIGN.LEFT
        run2 = p2.runs[0]
        run2.font.size = Pt(10)
        run2.font.color.rgb = hex_to_rgb("2B3C52")
    add_footer(slide)


def make_matrix(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, "F8FAFC")
    add_title(slide, "人口五维度 × 应急管理五模块对应关系", subtitle="人口要素贯穿应急管理全流程")
    rows = [
        ["人口规模", "风险评估、资源管理", "区域总人口决定灾害潜在损失规模与避难所容量测算"],
        ["人口结构", "应急预案、现场处置", "老龄化、儿童占比高，预案需优先考虑特殊人群救助"],
        ["人口空间分布", "风险识别、应急规划", "高密度区域风险更大；避难场所、疏散路线需适应分布特征"],
        ["人口迁移流动", "现场处置、灾后安置", "流动人口登记不全会增加预警通知和转移安置难度"],
        ["人口质量与健康", "公共卫生应急、灾后重建", "健康基础差异影响传染病传播、防疫和心理恢复"],
    ]
    table = slide.shapes.add_table(len(rows) + 1, 3, Inches(0.6), Inches(1.6), Inches(12.1), Inches(4.8)).table
    headers = ["社会学人口维度", "对应应急管理模块", "具体联系"]
    for col, text in enumerate(headers):
        cell = table.cell(0, col)
        cell.text = text
        for paragraph in cell.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.size = Pt(15)
                run.font.color.rgb = hex_to_rgb("FFFFFF")
        cell.fill.fore_color.rgb = hex_to_rgb("0B3D91")
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = val
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(11)
                    run.font.color.rgb = hex_to_rgb("1F2D3D")
            if r_idx % 2 == 0:
                cell.fill.fore_color.rgb = hex_to_rgb("F0F5FF")
            else:
                cell.fill.fore_color.rgb = hex_to_rgb("FFFFFF")
    add_footer(slide)


def make_cases(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, "F7F9FC")
    add_title(slide, "现实案例分析", subtitle="人口因素如何影响应急管理")
    left = 0.8
    top = 1.6
    width = 5.8
    height = 4.8
    # case 1
    panel1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    panel1.fill.solid(); panel1.fill.fore_color.rgb = hex_to_rgb("EAF2FF"); panel1.line.color.rgb = hex_to_rgb("9AB9D9")
    text1 = slide.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.2), Inches(width - 0.4), Inches(height - 0.4))
    tf1 = text1.text_frame
    p = tf1.paragraphs[0]; p.text = "案例一：洪涝灾害中的人口因素"; p.runs[0].font.bold = True; p.runs[0].font.size = Pt(20); p.runs[0].font.color.rgb = hex_to_rgb("0B3D91")
    for idx, line in enumerate([
        "• 社会学视角：灾区老龄化程度高，留守老人较多，外来务工人员比例大。",
        "• 应急管理措施：风险评估、老人上门救助、慢性病药品保障、人员摸排。",
        "• 启示：人口结构与流动特征直接指导救援和安置方案。",
    ]):
        p = tf1.add_paragraph(); p.text = line; p.bullet = True; p.level = 0; p.runs[0].font.size = Pt(13); p.runs[0].font.color.rgb = hex_to_rgb("1F2D3D")
    # case 2
    panel2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(top), Inches(width), Inches(height))
    panel2.fill.solid(); panel2.fill.fore_color.rgb = hex_to_rgb("FCEFE8"); panel2.line.color.rgb = hex_to_rgb("E7BEAB")
    text2 = slide.shapes.add_textbox(Inches(7.1), Inches(top + 0.2), Inches(width - 0.4), Inches(height - 0.4))
    tf2 = text2.text_frame
    p = tf2.paragraphs[0]; p.text = "案例二：城市高层建筑火灾"; p.runs[0].font.bold = True; p.runs[0].font.size = Pt(20); p.runs[0].font.color.rgb = hex_to_rgb("A0441B")
    for idx, line in enumerate([
        "• 社会学视角：城市人口高度密集，租住群体较多，信息获取不充分。",
        "• 应急管理措施：强化社区宣传、完善登记、优化疏散和搜救。",
        "• 启示：人口密度和流动性质决定灾害损失和救援难度。",
    ]):
        p = tf2.add_paragraph(); p.text = line; p.bullet = True; p.level = 0; p.runs[0].font.size = Pt(13); p.runs[0].font.color.rgb = hex_to_rgb("1F2D3D")
    add_footer(slide)


def make_summary(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, "F8FAFC")
    add_title(slide, "总结与专业思考", subtitle="从人口视角理解应急管理")
    bullets = [
        "社会学概论中的“人口”章节，是应急管理的重要社会理论基础。",
        "人口与应急管理在价值目标、研究对象、风险逻辑、治理思维上高度一致。",
        "应急管理不能只关注灾害技术本身，还必须考虑人口的社会属性与脆弱性。",
        "风险评估、预案编制和救援方案必须重视人口规模、结构、分布和流动特征。",
        "未来应急管理需要更加强调“以人为本”的社会治理思路。",
    ]
    add_bullets(slide, bullets, left=0.9, top=1.6, width=11.5, height=4.6, font_size=19)
    add_footer(slide)


def make_end(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, "F1F5FB")
    title = slide.shapes.add_textbox(Inches(0.8), Inches(2.3), Inches(11.5), Inches(1.0))
    tf = title.text_frame
    p = tf.paragraphs[0]
    p.text = "谢谢观看"
    p.alignment = PP_ALIGN.CENTER
    run = p.runs[0]
    run.font.size = Pt(32)
    run.font.bold = True
    run.font.color.rgb = hex_to_rgb("0B3D91")

    subtitle = slide.shapes.add_textbox(Inches(0.8), Inches(3.8), Inches(11.5), Inches(0.7))
    tf2 = subtitle.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = "欢迎老师和同学批评指正"
    p2.alignment = PP_ALIGN.CENTER
    run2 = p2.runs[0]
    run2.font.size = Pt(22)
    run2.font.color.rgb = hex_to_rgb("40536B")

    add_footer(slide, "Q&A")


def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    make_cover(prs)
    make_toc(prs)
    make_sociology_review(prs)
    make_emergency_framework(prs)
    make_common_logic(prs)
    make_matrix(prs)
    make_cases(prs)
    make_summary(prs)
    make_end(prs)
    prs.save(OUT_FILE)
    print(f"Generated: {OUT_FILE}")


if __name__ == "__main__":
    build_presentation()
